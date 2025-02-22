import docx
import pdfplumber
import csv
import pptxpip
import requests
from io import BytesIO
from bs4 import BeautifulSoup
import youtube_transcript_api
import gspread
from google.oauth2.service_account import Credentials as ServiceAccountCredentials
from fpdf import FPDF
from typing import Optional
import logging
from some_module import GoogleGenerativeAI, GoogleGenerativeAIEmbeddings, JsonOutputParser, NotesFormat, Chroma  # Replace 'some_module' with the actual module name

def read_text_file(filepath):
    with open(filepath, 'r', encoding='utf-8') as file:
        return file.read()

class NotesGenerator:
    def __init__(self, args=None, vectorstore_class=None, prompt=None, embedding_model=None, model=None, parser=None, verbose=False):
        default_config = {
            "model": GoogleGenerativeAI(model="gemini-1.5-flash"),
            "embedding_model": GoogleGenerativeAIEmbeddings(model='models/embedding-001'),
            "parser": JsonOutputParser(pydantic_object=NotesFormat),
            "prompt": read_text_file("prompt/notes-generator-prompt.txt"),
            "vectorstore_class": vectorstore_class or Chroma
        }

        self.prompt = prompt or default_config["prompt"]
        self.model = model or default_config["model"]
        self.parser = parser or default_config["parser"]
        self.embedding_model = embedding_model or default_config["embedding_model"]
        self.vectorstore_class = default_config["vectorstore_class"]
        self.vectorstore, self.retriever, self.runner = None, None, None
        self.args = args
        self.verbose = verbose

        if not self.vectorstore_class:
            raise ValueError("Vectorstore must be provided")
        if not args or not args.topic:
            raise ValueError("Topic must be provided")
        if not args.format_type:
            raise ValueError("Format type must be provided (bullet, paragraph, table)")
        if args.format_type not in ["bullet", "paragraph", "table"]:
            raise ValueError("Invalid format type. Choose from 'bullet', 'paragraph', or 'table'.")

    def extract_text_from_pdf(self, file):
        with pdfplumber.open(file) as pdf:
            return "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])

    def extract_text_from_docx(self, file):
        doc = docx.Document(file)
        return "\n".join([para.text for para in doc.paragraphs])

    def extract_text_from_csv(self, file):
        content = file.read().decode("utf-8").splitlines()
        reader = csv.reader(content)
        return "\n".join([", ".join(row) for row in reader])

    def extract_text_from_ppt(self, file):
        presentation = pptx.Presentation(file)
        return "\n".join([slide.text for slide in presentation.slides if slide.text])

    def extract_text_from_plain_text(self, file):
        return file.read().decode("utf-8")

    def extract_text_from_website(self, url):
        response = requests.get(url)
        soup = BeautifulSoup(response.text, 'html.parser')
        return "\n".join([p.text for p in soup.find_all('p')])

    def extract_text_from_youtube(self, video_id):
        transcript = youtube_transcript_api.YouTubeTranscriptApi.get_transcript(video_id)
        return "\n".join([entry['text'] for entry in transcript])

    def extract_text_from_google_sheets(self, sheet_url, credentials_json):
        scope = ["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]
        creds = ServiceAccountCredentials.from_json_keyfile_name(credentials_json, scope)
        client = gspread.authorize(creds)
        sheet = client.open_by_url(sheet_url).sheet1
        return "\n".join([", ".join(row) for row in sheet.get_all_values()])

    def generate_notes(self, content: str, format_type: str = "bullet"): 
        if format_type == "bullet":
            return f"Summary:\n- " + "\n- ".join(content.split(". "))
        elif format_type == "paragraph":
            return f"Summary:\n" + content.replace(". ", ".\n")
        elif format_type == "table":
            rows = content.split(". ")
            return "\n".join([f"| {row} |" for row in rows])
        else:
            return "Invalid format type."

    def export_to_docx(self, notes: str, filename: str):
        doc = docx.Document()
        doc.add_paragraph(notes)
        doc.save(filename)

    def export_to_pdf(self, notes: str, filename: str):
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, notes)
        pdf.output(filename)

    def export_to_txt(self, notes: str, filename: str):
        with open(filename, "w", encoding="utf-8") as file:
            file.write(notes)
